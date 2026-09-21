import openpyxl, json
wb = openpyxl.load_workbook("/tmp/ali.xlsx", read_only=True, data_only=True)
for ws in wb.worksheets:
    print("="*70)
    print("SHEET:", ws.title, "dims:", ws.max_row, "x", ws.max_column)
    rows = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
    if rows:
        print("HEADER:", rows[0])
        for r in rows[1:]:
            print("  ROW :", r)
